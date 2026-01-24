from typing import List, Dict
from decimal import Decimal
from datetime import date
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from loguru import logger


class PensionReportService:
    
    @staticmethod
    def generate_pension_declaration_excel(
        organization_name: str,
        tin_number: str,
        period: str,
        payroll_items: List[Dict]
    ) -> bytes:
        """
        Generate Ethiopian Pension Declaration Form (Excel Format)
        Required for submission to Ethiopian Social Security Authority
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Pension Declaration"
        
        ws.merge_cells('A1:H1')
        title_cell = ws['A1']
        title_cell.value = "ETHIOPIAN SOCIAL SECURITY AUTHORITY - PENSION CONTRIBUTION DECLARATION"
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        title_cell.fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")
        title_cell.font = Font(bold=True, color="FFFFFF", size=14)
        ws.row_dimensions[1].height = 25
        
        ws['A3'] = "Organization Name:"
        ws['B3'] = organization_name
        ws['A4'] = "TIN Number:"
        ws['B4'] = tin_number
        ws['A5'] = "Period:"
        ws['B5'] = period
        ws['A6'] = "Submission Date:"
        ws['B6'] = date.today().strftime('%d/%m/%Y')
        
        for row in range(3, 7):
            ws.cell(row=row, column=1).font = Font(bold=True)
        
        header_row = 8
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = [
            "S/N",
            "Employee ID",
            "Full Name",
            "Gross Salary (ETB)",
            "Employee Pension (7%)",
            "Employer Pension (11%)",
            "Total Pension (18%)",
            "Remarks"
        ]
        
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        
        ws.row_dimensions[header_row].height = 30
        
        total_gross = Decimal("0")
        total_employee_pension = Decimal("0")
        total_employer_pension = Decimal("0")
        total_pension = Decimal("0")
        
        data_start_row = header_row + 1
        
        for idx, item in enumerate(payroll_items, start=data_start_row):
            gross = Decimal(str(item.get('gross_salary', 0)))
            emp_pension = Decimal(str(item.get('employee_pension', 0)))
            empr_pension = Decimal(str(item.get('employer_pension', 0)))
            total_pens = emp_pension + empr_pension
            
            total_gross += gross
            total_employee_pension += emp_pension
            total_employer_pension += empr_pension
            total_pension += total_pens
            
            ws.cell(row=idx, column=1, value=idx - data_start_row + 1).border = border
            ws.cell(row=idx, column=2, value=item.get('employee_id', '')).border = border
            ws.cell(row=idx, column=3, value=item.get('worker_name', '')).border = border
            
            gross_cell = ws.cell(row=idx, column=4, value=float(gross))
            gross_cell.number_format = '#,##0.00'
            gross_cell.border = border
            
            emp_pens_cell = ws.cell(row=idx, column=5, value=float(emp_pension))
            emp_pens_cell.number_format = '#,##0.00'
            emp_pens_cell.border = border
            
            empr_pens_cell = ws.cell(row=idx, column=6, value=float(empr_pension))
            empr_pens_cell.number_format = '#,##0.00'
            empr_pens_cell.border = border
            
            total_pens_cell = ws.cell(row=idx, column=7, value=float(total_pens))
            total_pens_cell.number_format = '#,##0.00'
            total_pens_cell.border = border
            
            tax_rule = item.get('tax_rule_applied', '')
            remarks = "Regular" if tax_rule == "standard_employee" else "Excluded"
            ws.cell(row=idx, column=8, value=remarks).border = border
        
        total_row = len(payroll_items) + data_start_row
        ws.merge_cells(f'A{total_row}:C{total_row}')
        total_label = ws.cell(row=total_row, column=1, value="TOTAL")
        total_label.font = Font(bold=True, size=12)
        total_label.alignment = Alignment(horizontal='right')
        total_label.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        total_label.border = border
        
        for col, value in enumerate([total_gross, total_employee_pension, total_employer_pension, total_pension], start=4):
            cell = ws.cell(row=total_row, column=col, value=float(value))
            cell.number_format = '#,##0.00'
            cell.font = Font(bold=True, size=12)
            cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            cell.border = border
        
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 18
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 18
        ws.column_dimensions['G'].width = 18
        ws.column_dimensions['H'].width = 15
        
        signature_row = total_row + 3
        ws.cell(row=signature_row, column=1, value="Prepared By:")
        ws.cell(row=signature_row + 1, column=1, value="Name: _____________________")
        ws.cell(row=signature_row + 2, column=1, value="Signature: _________________")
        ws.cell(row=signature_row + 3, column=1, value="Date: _____________________")
        
        ws.cell(row=signature_row, column=5, value="Approved By:")
        ws.cell(row=signature_row + 1, column=5, value="Name: _____________________")
        ws.cell(row=signature_row + 2, column=5, value="Signature: _________________")
        ws.cell(row=signature_row + 3, column=5, value="Date: _____________________")
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        logger.info(f"Pension declaration Excel generated for {len(payroll_items)} workers")
        return output.getvalue()
    
    @staticmethod
    def generate_pension_declaration_pdf(
        organization_name: str,
        tin_number: str,
        period: str,
        payroll_items: List[Dict]
    ) -> bytes:
        """
        Generate Ethiopian Pension Declaration Form (PDF Format)
        Professional PDF report for government submission
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5*inch, bottomMargin=0.5*inch)
        
        story = []
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=14,
            textColor=colors.HexColor('#0070C0'),
            spaceAfter=12,
            alignment=1
        )
        
        title = Paragraph("ETHIOPIAN SOCIAL SECURITY AUTHORITY<br/>PENSION CONTRIBUTION DECLARATION", title_style)
        story.append(title)
        story.append(Spacer(1, 0.2*inch))
        
        info_data = [
            ["Organization Name:", organization_name, "Period:", period],
            ["TIN Number:", tin_number, "Submission Date:", date.today().strftime('%d/%m/%Y')]
        ]
        
        info_table = Table(info_data, colWidths=[1.5*inch, 2.5*inch, 1.2*inch, 1.8*inch])
        info_table.setStyle(TableStyle([
            ('FONT', (0, 0), (0, -1), 'Helvetica-Bold', 10),
            ('FONT', (2, 0), (2, -1), 'Helvetica-Bold', 10),
            ('FONT', (1, 0), (1, -1), 'Helvetica', 10),
            ('FONT', (3, 0), (3, -1), 'Helvetica', 10),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 0.3*inch))
        
        table_data = [
            ["S/N", "Emp ID", "Full Name", "Gross\n(ETB)", "Employee\n7% (ETB)", "Employer\n11% (ETB)", "Total\n18% (ETB)"]
        ]
        
        total_gross = Decimal("0")
        total_employee_pension = Decimal("0")
        total_employer_pension = Decimal("0")
        
        for idx, item in enumerate(payroll_items, start=1):
            gross = Decimal(str(item.get('gross_salary', 0)))
            emp_pension = Decimal(str(item.get('employee_pension', 0)))
            empr_pension = Decimal(str(item.get('employer_pension', 0)))
            total_pens = emp_pension + empr_pension
            
            total_gross += gross
            total_employee_pension += emp_pension
            total_employer_pension += empr_pension
            
            table_data.append([
                str(idx),
                item.get('employee_id', ''),
                item.get('worker_name', ''),
                f"{float(gross):,.2f}",
                f"{float(emp_pension):,.2f}",
                f"{float(empr_pension):,.2f}",
                f"{float(total_pens):,.2f}"
            ])
        
        total_pension = total_employee_pension + total_employer_pension
        table_data.append([
            "", "", "TOTAL",
            f"{float(total_gross):,.2f}",
            f"{float(total_employee_pension):,.2f}",
            f"{float(total_employer_pension):,.2f}",
            f"{float(total_pension):,.2f}"
        ])
        
        pension_table = Table(table_data, colWidths=[0.5*inch, 0.8*inch, 2*inch, 1*inch, 1*inch, 1*inch, 1*inch])
        pension_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#D9E1F2')),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
        ]))
        
        story.append(pension_table)
        story.append(Spacer(1, 0.4*inch))
        
        signature_data = [
            ["Prepared By: _____________________", "Approved By: _____________________"],
            ["Date: _____________________", "Date: _____________________"]
        ]
        
        signature_table = Table(signature_data, colWidths=[3.5*inch, 3.5*inch])
        signature_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
        ]))
        story.append(signature_table)
        
        doc.build(story)
        
        buffer.seek(0)
        logger.info(f"Pension declaration PDF generated for {len(payroll_items)} workers")
        return buffer.getvalue()
